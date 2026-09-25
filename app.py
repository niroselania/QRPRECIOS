import io
import os
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime

import pandas as pd
import openpyxl
from PIL import Image, ImageDraw, ImageFont
from flask import Flask, jsonify, request, render_template, send_file, flash, redirect, url_for
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "equivalencia-app-secret")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR") or os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)

GRANDE_CACHE_PATH_XLS = os.path.join(DATA_DIR, "planilla_grande.xls")
GRANDE_CACHE_PATH_XLSX = os.path.join(DATA_DIR, "planilla_grande.xlsx")
GRANDE_CACHE_META = os.path.join(DATA_DIR, "planilla_grande.meta")
STOCK_CACHE_PATH = os.path.join(DATA_DIR, "stock_precios.xlsx")
STOCK_CACHE_META = os.path.join(DATA_DIR, "stock_precios.meta")

ALLOWED_EXT = {".xls", ".xlsx"}
FONT_BOLD = os.path.join(BASE_DIR, "fonts", "DejaVuSans-Bold.ttf")
FONT_REGULAR = os.path.join(BASE_DIR, "fonts", "DejaVuSans.ttf")
IVA = 1.21


def _ext(filename):
    return os.path.splitext(filename)[1].lower()


def _encabezado(valor):
    texto = unicodedata.normalize("NFKD", str(valor).strip()).encode("ascii", "ignore").decode()
    return " ".join(texto.upper().split())


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    return texto[:-2] if texto.endswith(".0") else texto


def _clave_sku(sku):
    return _texto(sku).upper().replace(" ", "")


def _clave_producto(sku, color):
    return f"{_clave_sku(sku)}{_texto(color)}".upper().replace(" ", "")


def _clave_barra(codigo):
    texto = _texto(codigo).replace(" ", "")
    return str(int(texto)) if texto.isdigit() else texto.upper()


def _precio(valor):
    texto = _texto(valor)
    if not texto or texto == "-":
        return None
    if isinstance(valor, (int, float)):
        return round(float(valor))
    limpio = texto.replace("$", "").replace(" ", "")
    if "," in limpio and "." in limpio:
        limpio = limpio.replace(".", "").replace(",", ".")
    elif "," in limpio:
        limpio = limpio.replace(",", ".")
    return round(float(limpio))


def _buscar_columna(headers, *nombres_posibles):
    for nombre in nombres_posibles:
        if nombre in headers:
            return headers[nombre]
    return None


def _read_grande(path):
    """Arma un índice código de barras -> SKU, color, talle y descripción."""
    ext = _ext(path)
    engine = "xlrd" if ext == ".xls" else "openpyxl"
    xl = pd.ExcelFile(path, engine=engine)

    index = {}
    for sheet in xl.sheet_names:
        df = xl.parse(sheet, header=None, dtype=str)
        if df.empty:
            continue
        encabezados = {
            _encabezado(valor): columna
            for columna, valor in enumerate(df.iloc[0])
            if not pd.isna(valor) and str(valor).strip()
        }
        col_codigo = _buscar_columna(
            encabezados, "CODIGO", "CODIGOS", "CODIGO DE BARRA", "CODIGOS DE BARRA",
            "BARRA", "BARRAS", "BARR", "BARCODE",
        )
        col_articulo = _buscar_columna(encabezados, "ARTICULO", "ARTICULOS", "ARTIC", "SKU COLOR", "CONCAT")
        col_sku = _buscar_columna(encabezados, "SKU")
        col_color = _buscar_columna(encabezados, "COLOR")
        col_talle = _buscar_columna(encabezados, "TALLE")
        col_descripcion = _buscar_columna(encabezados, "DESCRIPCION", "NOMBRE", "PRODUCTO")

        tiene_encabezados = col_codigo is not None and (
            col_articulo is not None or (col_sku is not None and col_color is not None)
        )
        if tiene_encabezados:
            df = df.iloc[1:]
        else:
            col_codigo, col_articulo, col_sku, col_color = 0, 1, None, None
            col_descripcion, col_talle = 2, 5

        for _, row in df.iterrows():
            codigo = row.get(col_codigo)
            if pd.isna(codigo):
                continue
            codigo = _texto(codigo)
            if not codigo:
                continue

            sku = _texto(row.get(col_sku)) if col_sku is not None else ""
            color = _texto(row.get(col_color)) if col_color is not None else ""
            if (not sku or not color) and col_articulo is not None and not pd.isna(row.get(col_articulo)):
                partes = str(row.get(col_articulo)).strip().split(maxsplit=1)
                if len(partes) == 2:
                    sku, color = partes
            if not sku or not color:
                continue

            descripcion = _texto(row.get(col_descripcion)) if col_descripcion is not None else ""
            talle = _texto(row.get(col_talle)) if col_talle is not None else ""
            index.setdefault(
                _clave_barra(codigo),
                {"codigo": codigo, "sku": sku, "color": color, "talle": talle, "descripcion": descripcion},
            )
    return index


_BARCODE_CACHE = {"path": None, "mtime": None, "index": None}


def _get_barcode_index(path):
    global _BARCODE_CACHE
    mtime = os.path.getmtime(path)
    if _BARCODE_CACHE["path"] == path and _BARCODE_CACHE["mtime"] == mtime:
        return _BARCODE_CACHE["index"]
    index = _read_grande(path)
    _BARCODE_CACHE = {"path": path, "mtime": mtime, "index": index}
    return index


_PRICE_CACHE = {"path": None, "mtime": None, "prices": None, "by_sku": None}


def _get_price_lookup(path):
    """Lee RETAIL del STOCK (precio sin IVA) agrupado por SKU + color y también por SKU."""
    global _PRICE_CACHE
    mtime = os.path.getmtime(path)
    if _PRICE_CACHE["path"] == path and _PRICE_CACHE["mtime"] == mtime:
        return _PRICE_CACHE["prices"], _PRICE_CACHE["by_sku"]

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hoja = workbook["STOCK"] if "STOCK" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    filas = hoja.iter_rows(values_only=True)
    try:
        encabezado = next(filas)
    except StopIteration:
        workbook.close()
        raise ValueError("La planilla de stock está vacía.")

    headers = {_encabezado(valor): i for i, valor in enumerate(encabezado) if valor is not None}
    col_sku = _buscar_columna(headers, "SKU", "ARTICULO", "ARTICULOS")
    col_color = _buscar_columna(headers, "COLOR")
    col_precio = _buscar_columna(
        headers, "RETAIL", "PRECIO SIN IVA", "PRECIO SIN IVA.", "P RETAIL", "PRECIO", "PVP"
    )
    if col_sku is None:
        col_sku = 1
    if col_color is None:
        col_color = 2
    if col_precio is None:
        col_precio = 10

    values_by_product = defaultdict(list)
    for row in filas:
        if row is None or len(row) <= max(col_sku, col_color, col_precio):
            continue
        sku_key = _clave_sku(row[col_sku])
        color = _texto(row[col_color])
        amount = _precio(row[col_precio])
        if sku_key and amount is not None:
            values_by_product[(sku_key, color)].append(amount)
    workbook.close()

    prices = {}
    by_sku = defaultdict(list)
    for (sku_key, color), values in values_by_product.items():
        amount = Counter(values).most_common(1)[0][0]
        prices[_clave_producto(sku_key, color)] = amount
        by_sku[sku_key].append({"sku": sku_key, "color": color, "precio_sin_iva": amount})
    _PRICE_CACHE = {"path": path, "mtime": mtime, "prices": prices, "by_sku": dict(by_sku)}
    return prices, by_sku


def _resolver_grande_cache():
    if os.path.exists(GRANDE_CACHE_PATH_XLS):
        return GRANDE_CACHE_PATH_XLS
    if os.path.exists(GRANDE_CACHE_PATH_XLSX):
        return GRANDE_CACHE_PATH_XLSX
    return None


def _resolver_stock_cache():
    return STOCK_CACHE_PATH if os.path.exists(STOCK_CACHE_PATH) else None


def _leer_meta(path):
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return f.read().strip()


def _guardar_meta(path, filename):
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"{secure_filename(filename)} - cargada {datetime.now().strftime('%d/%m/%Y %H:%M')}")


def _money(amount):
    return "$ " + f"{round(amount):,}".replace(",", ".")


def _fuente_que_entra(path, size_px, texto, max_width):
    font = ImageFont.truetype(path, size=size_px)
    probe = ImageDraw.Draw(Image.new("L", (1, 1)))
    while size_px > 8:
        font = ImageFont.truetype(path, size=size_px)
        ancho = probe.textbbox((0, 0), texto, font=font)[2]
        if ancho <= max_width:
            return font
        size_px -= 1
    return font


def _texto_centrado(draw, y, texto, font, ancho):
    bbox = draw.textbbox((0, 0), texto, font=font)
    x = max(0, (ancho - (bbox[2] - bbox[0])) // 2)
    draw.text((x, y), texto, font=font, fill=0)
    return bbox[3] - bbox[1]


def generar_etiqueta(precio_sin_iva, dpi=300):
    """Etiqueta vertical 15 x 30 mm solo con precio sin IVA y con IVA."""
    px_mm = dpi / 25.4
    W, H = round(15.0 * px_mm), round(30.0 * px_mm)
    img = Image.new("L", (W, H), color=255)
    draw = ImageDraw.Draw(img)
    max_width = W - round(1.2 * px_mm)

    sin_iva = _money(precio_sin_iva)
    con_iva = _money(round(precio_sin_iva * IVA))
    font_label = _fuente_que_entra(FONT_REGULAR, round(2.1 * px_mm), "SIN IVA", max_width)
    font_price = _fuente_que_entra(
        FONT_BOLD, round(3.4 * px_mm), sin_iva if len(sin_iva) >= len(con_iva) else con_iva, max_width
    )

    bloques = [
        ("SIN IVA", font_label),
        (sin_iva, font_price),
        ("CON IVA", font_label),
        (con_iva, font_price),
    ]
    altos = []
    probe = ImageDraw.Draw(Image.new("L", (1, 1)))
    for texto, font in bloques:
        bbox = probe.textbbox((0, 0), texto, font=font)
        altos.append(bbox[3] - bbox[1])
    hueco = round(1.3 * px_mm)
    total = sum(altos) + hueco * 3
    y = max(round(1.2 * px_mm), (H - total) // 2)
    for (texto, font), alto in zip(bloques, altos):
        _texto_centrado(draw, y, texto, font, W)
        y += alto + hueco

    salida = io.BytesIO()
    img.save(salida, format="PNG", dpi=(dpi, dpi))
    salida.seek(0)
    return salida


def _error_etiqueta(mensaje, status=400):
    if request.headers.get("X-Requested-With") == "fetch" or "application/json" in request.headers.get("Accept", ""):
        return jsonify({"error": mensaje}), status
    flash(mensaje)
    return redirect(url_for("index"))


@app.route("/etiqueta")
def etiqueta():
    return redirect(url_for("index"))


def _resolver_precio(datos):
    barcode_scanned = _texto(datos.get("barcode") or "").replace("\r", "").replace("\n", "").replace("\t", "")
    sku = _clave_sku(datos.get("sku") or "")
    color = _texto(datos.get("color") or "")

    stock_path = _resolver_stock_cache()
    if not stock_path:
        return None, _error_etiqueta("No hay planilla de STOCK con precios cargada. Subila primero.")

    try:
        prices, by_sku = _get_price_lookup(stock_path)
    except Exception as e:
        return None, _error_etiqueta(f"Error leyendo los precios: {e}", 500)

    if barcode_scanned:
        grande_path = _resolver_grande_cache()
        if not grande_path:
            return None, _error_etiqueta("No hay planilla de equivalencia cargada. Subila primero o buscá por SKU.")
        try:
            producto = _get_barcode_index(grande_path).get(_clave_barra(barcode_scanned))
        except Exception as e:
            return None, _error_etiqueta(f"Error buscando el producto: {e}", 500)
        if not producto:
            return None, _error_etiqueta(
                f"No encontré el código {barcode_scanned}. Probá buscarlo por SKU."
            )
        sku = _clave_sku(producto["sku"])
        color = _texto(producto["color"])
        precio_sin_iva = prices.get(_clave_producto(sku, color))
        if precio_sin_iva is None:
            return None, _error_etiqueta(f"No encontré precio para {sku} {color} en la planilla STOCK.")
        return {"sku": sku, "color": color, "precio_sin_iva": precio_sin_iva}, None

    if not sku:
        return None, _error_etiqueta("Escaneá un código de barras o ingresá el SKU.")

    opciones = list(by_sku.get(sku) or [])
    if not opciones:
        return None, _error_etiqueta(f"No encontré el SKU {sku} en la planilla STOCK.")

    if color:
        coincidencias = [op for op in opciones if _texto(op["color"]).upper() == color.upper()]
        if not coincidencias:
            return None, _error_etiqueta(f"No encontré el SKU {sku} en color {color}.")
        opciones = coincidencias

    precios = {op["precio_sin_iva"] for op in opciones}
    if len(opciones) == 1 or len(precios) == 1:
        return opciones[0], None

    return None, jsonify({"opciones": opciones})


@app.route("/etiqueta/generar", methods=["GET", "POST"])
def etiqueta_generar():
    producto, error = _resolver_precio(request.values)
    if error is not None:
        return error

    try:
        imagen = generar_etiqueta(producto["precio_sin_iva"])
    except Exception as e:
        return _error_etiqueta(f"No pude generar la etiqueta: {e}", 500)

    descargar = request.values.get("descargar") == "1"
    nombre = f"etiqueta_{secure_filename(producto['sku'])}_{secure_filename(producto['color'] or 'precio')}.png"
    return send_file(
        imagen,
        mimetype="image/png",
        as_attachment=descargar,
        download_name=nombre,
    )


@app.route("/", methods=["GET"])
def index():
    return render_template(
        "index.html",
        grande_cached=_resolver_grande_cache() is not None,
        grande_info=_leer_meta(GRANDE_CACHE_META),
        stock_cached=_resolver_stock_cache() is not None,
        stock_info=_leer_meta(STOCK_CACHE_META),
    )


@app.route("/subir-equivalencia", methods=["POST"])
def subir_equivalencia():
    grande_file = request.files.get("grande")
    if not grande_file or grande_file.filename == "":
        flash("Subí la planilla de equivalencia.")
        return redirect(url_for("index"))
    if _ext(grande_file.filename) not in ALLOWED_EXT:
        flash("La planilla de equivalencia debe ser .xls o .xlsx.")
        return redirect(url_for("index"))

    grande_ext = _ext(grande_file.filename)
    grande_path = GRANDE_CACHE_PATH_XLS if grande_ext == ".xls" else GRANDE_CACHE_PATH_XLSX
    other_path = GRANDE_CACHE_PATH_XLSX if grande_ext == ".xls" else GRANDE_CACHE_PATH_XLS
    if os.path.exists(other_path):
        os.remove(other_path)
    grande_file.save(grande_path)

    global _BARCODE_CACHE
    _BARCODE_CACHE = {"path": None, "mtime": None, "index": None}
    try:
        _get_barcode_index(grande_path)
    except Exception as e:
        flash(f"La planilla se guardó, pero no pude leerla: {e}")
        return redirect(url_for("index"))

    _guardar_meta(GRANDE_CACHE_META, grande_file.filename)
    flash("Planilla de equivalencia cargada correctamente.")
    return redirect(url_for("index"))


@app.route("/subir-stock", methods=["POST"])
def subir_stock():
    stock_file = request.files.get("stock")
    if not stock_file or stock_file.filename == "":
        flash("Subí la planilla STOCK con los precios.")
        return redirect(url_for("index"))
    if _ext(stock_file.filename) != ".xlsx":
        flash("La planilla STOCK debe ser un archivo .xlsx.")
        return redirect(url_for("index"))

    stock_file.save(STOCK_CACHE_PATH)
    global _PRICE_CACHE
    _PRICE_CACHE = {"path": None, "mtime": None, "prices": None, "by_sku": None}
    try:
        _get_price_lookup(STOCK_CACHE_PATH)
    except Exception as e:
        flash(f"La planilla se guardó, pero no pude leer los precios: {e}")
        return redirect(url_for("index"))

    _guardar_meta(STOCK_CACHE_META, stock_file.filename)
    flash("Planilla STOCK con precios cargada correctamente.")
    return redirect(url_for("index"))


@app.route("/health")
def health():
    return {"status": "ok"}


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8000)), threaded=True)
